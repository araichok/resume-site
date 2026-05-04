from flask import Flask, render_template, request, session, redirect, url_for
from werkzeug.utils import secure_filename

from docx import Document
from flask import send_file

import os

app = Flask(__name__)

app.secret_key = "resume_secret_key"

UPLOAD_FOLDER = "static/uploads"
app.config["UPLOAD_FOLDER"] = UPLOAD_FOLDER

if not os.path.exists(UPLOAD_FOLDER):
    os.makedirs(UPLOAD_FOLDER)


@app.route("/")
def home():
    lang = request.args.get("lang")

    if lang:
        session["lang"] = lang

    lang = session.get("lang", "ru")

    documents = [
        {
            "title_ru": "Титульный лист отчета",
            "title_kz": "Есептің титул парағы",
            "file": "1. Титульный лист отчета.docx"
        },
        {
            "title_ru": "Трехсторонний договор",
            "title_kz": "Үшжақты келісім-шарт",
            "file": "2. Трехсторонний договор на проведение профессиональной практики.docx"
        },
        {
            "title_ru": "Направление",
            "title_kz": "Жолдама",
            "file": "3. Направление.docx"
        },
        {
            "title_ru": "Рабочий план-график",
            "title_kz": "Жұмыс жоспары",
            "file": "4. Рабочий план-график.docx"
        },
        {
            "title_ru": "Дневник-отчет о прохождении практики",
            "title_kz": "Күнделік-тәжірибеден өту туралы есеп",
            "file": "5. Дневник-отчет о прохождении практики.docx"
        },
        {
            "title_ru": "Характеристика на обучающегося",
            "title_kz": "Білім алушыға мінездеме",
            "file": "6. Характеристика на обучающегося.docx"
        }
    ]

    return render_template("home.html", documents=documents, lang=lang)


@app.route("/builder", methods=["GET", "POST"])
def builder():
    lang = request.args.get("lang")

    if lang:
        session["lang"] = lang

    lang = session.get("lang", "ru")

    if request.method == "POST":
        photo = request.files.get("photo")
        filename = None

        if photo and photo.filename != "":
            filename = secure_filename(photo.filename)
            photo.save(os.path.join(app.config["UPLOAD_FOLDER"], filename))

        data = {
            "name": request.form.get("name"),
            "phone": request.form.get("phone"),
            "email": request.form.get("email"),
            "city": request.form.get("city"),
            "about": request.form.get("about"),
            "education": request.form.get("education"),
            "experience": request.form.get("experience"),
            "skills": request.form.get("skills"),
            "sport_achievements": request.form.get("sport_achievements"),
            "section_order": request.form.get("section_order", "about,education,experience,sport_achievements"),
            "template": request.form.get("template", "template1"),
            "photo": filename
        }

        session["resume"] = data
        return redirect(url_for("resume"))

    return render_template("builder.html", lang=lang)


def analyze_resume(data, lang):
    score = 0
    recommendations = []

    fields_ru = {
        "about": "О себе",
        "education": "Образование",
        "experience": "Опыт работы",
        "skills": "Навыки",
        "sport_achievements": "Спортивные достижения"
    }

    fields_kz = {
        "about": "Өзі туралы",
        "education": "Білімі",
        "experience": "Жұмыс тәжірибесі",
        "skills": "Дағдылар",
        "sport_achievements": "Спорттық жетістіктер"
    }

    fields = fields_ru if lang == "ru" else fields_kz

    for key, title in fields.items():
        value = data.get(key, "")

        if value and value.strip():
            score += 20
        else:
            if lang == "ru":
                recommendations.append(f"Добавьте раздел: {title}")
            else:
                recommendations.append(f"Бөлімді қосыңыз: {title}")

    about = data.get("about", "")
    skills = data.get("skills", "")
    experience = data.get("experience", "")

    if skills and len(skills.strip()) < 20:
        if lang == "ru":
            recommendations.append("Раздел «Навыки» можно сделать подробнее.")
        else:
            recommendations.append("«Дағдылар» бөлімін толығырақ жазуға болады.")

    if about and len(about.strip()) < 40:
        if lang == "ru":
            recommendations.append("Раздел «О себе» лучше написать более подробно.")
        else:
            recommendations.append("«Өзі туралы» бөлімін толығырақ жазған дұрыс.")

    if experience and len(experience.strip()) < 30:
        if lang == "ru":
            recommendations.append("Опыт работы или практику лучше описать подробнее.")
        else:
            recommendations.append("Жұмыс тәжірибесі немесе практиканы толығырақ сипаттаған дұрыс.")

    if score >= 80:
        level = "Хорошее резюме" if lang == "ru" else "Жақсы резюме"
    elif score >= 50:
        level = "Среднее резюме" if lang == "ru" else "Орташа резюме"
    else:
        level = "Резюме нужно доработать" if lang == "ru" else "Резюмені толықтыру қажет"

    return {
        "score": score,
        "level": level,
        "recommendations": recommendations
    }


@app.route("/resume")
def resume():
    data = session.get("resume")

    if not data:
        return redirect(url_for("builder"))

    lang = session.get("lang", "ru")

    analysis = analyze_resume(data, lang)

    return render_template(
        "resume.html",
        data=data,
        lang=lang,
        analysis=analysis
    )



@app.route("/download_word")
def download_word():
    data = session.get("resume")

    if not data:
        return redirect(url_for("builder"))

    doc = Document()

    # Заголовок
    doc.add_heading(data.get("name", "Резюме"), 0)

    # Контакты
    doc.add_paragraph(f"Телефон: {data.get('phone', '')}")
    doc.add_paragraph(f"Email: {data.get('email', '')}")
    doc.add_paragraph(f"Город: {data.get('city', '')}")

    doc.add_paragraph("")

    # Разделы
    def add_section(title, content):
        if content:
            doc.add_heading(title, level=1)
            doc.add_paragraph(content)

    add_section("О себе", data.get("about"))
    add_section("Образование", data.get("education"))
    add_section("Опыт работы", data.get("experience"))
    add_section("Навыки", data.get("skills"))
    add_section("Спортивные достижения", data.get("sport_achievements"))

    # Сохранение
    file_path = "resume.docx"
    doc.save(file_path)

    return send_file(file_path, as_attachment=True)


if __name__ == "__main__":
    app.run()